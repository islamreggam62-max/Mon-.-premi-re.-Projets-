#!/usr/bin/env python3
"""Courriers au format du Dépt (Times New Roman 12) sur papier à en-tête, à partir d'un JSON.

Usage :
    python3 lettre_tnr.py lettres.json --modele entete.docx [-o dossier]

`lettres.json` : un courrier (objet) ou plusieurs (liste). Champs :
  fichier      nom du .docx produit
  entete       lignes de l'expéditeur (la 1re porte « Dar El Beida, le » au taquet)
  destinataire lignes du destinataire (retrait 4253 twips)
  objet, ref   texte des lignes « Objet : » et « Réf : » (ref facultatif)
  appel        « Monsieur, » (facultatif)
  corps        paragraphes ; une ligne commençant par « - » devient un tiret
  cloture      formule de politesse ; signature : lignes (retrait 5103 twips)
  interligne   1.5 par défaut (1.3 pour tenir sur une page)
Les passages entre crochets [ … ] sont surlignés en jaune : ce sont les champs à compléter.
Exemple : modeles/resiliation.json (note M9 et lettre M10 bis).
"""
import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

F, S = "Times New Roman", Pt(12)
TAB, SIGN, ALINEA = Twips(4253), Twips(5103), Twips(709)
CROCHETS = re.compile(r"(\[[^\]]*\])")
NB = " "


def run(p, texte, gras=False, souligne=False):
    for morceau in CROCHETS.split(texte):
        if not morceau:
            continue
        r = p.add_run(morceau)
        r.font.name, r.font.size = F, S
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), F)
        r.bold, r.underline = gras, souligne
        if CROCHETS.fullmatch(morceau):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def para(doc, texte="", gras=False, gauche=None, alinea=None, apres=0, ligne=1.0,
         align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(apres), ligne
    p.alignment = align
    if gauche is not None:
        pf.left_indent = gauche
    if alinea is not None:
        pf.first_line_indent = alinea
    if texte:
        run(p, texte, gras)
    return p


def document(modele):
    doc = Document(modele) if modele else Document()
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    doc.styles["Normal"].font.name = F
    doc.styles["Normal"].font.size = S
    return doc


def courrier(doc, c):
    ligne = c.get("interligne", 1.5)
    entete = c["entete"]
    p = para(doc)
    p.paragraph_format.tab_stops.add_tab_stop(TAB, WD_TAB_ALIGNMENT.LEFT)
    run(p, entete[0] + "\tDar El Beida, le", gras=True)
    for l in entete[1:]:
        para(doc, l, gras=True)
    para(doc)
    para(doc)
    for l in c["destinataire"]:
        para(doc, l, gras=True, gauche=TAB)
    para(doc)
    para(doc)
    p = para(doc)
    run(p, "Objet" + NB + ":", gras=True, souligne=True)
    run(p, " " + c["objet"], gras=True)
    if c.get("ref"):
        p = para(doc)
        run(p, "Réf" + NB + ":", gras=True, souligne=True)
        run(p, " " + c["ref"], gras=True)
    para(doc)
    if c.get("appel"):
        para(doc, c["appel"], alinea=ALINEA, apres=6)
    for l in c.get("corps", []):
        if l.startswith("- "):
            para(doc, l, gauche=ALINEA, apres=4, ligne=ligne, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        else:
            para(doc, l, alinea=ALINEA, apres=6, ligne=ligne, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(doc, c.get("cloture", "Salutations."), alinea=ALINEA, ligne=ligne,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para(doc)
    for l in c.get("signature", ["LA DIRECTRICE JURIDIQUE", "M. LAMANI"]):
        para(doc, l, gras=True, gauche=SIGN)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("json")
    ap.add_argument("--modele", help=".docx dont l'en-tête (logo) et le pied de page sont repris")
    ap.add_argument("-o", "--dossier", default=".")
    args = ap.parse_args()
    data = json.loads(Path(args.json).read_text(encoding="utf-8"))
    for c in data if isinstance(data, list) else [data]:
        doc = document(args.modele)
        courrier(doc, c)
        sortie = Path(args.dossier) / c["fichier"]
        doc.save(sortie)
        print(sortie)


if __name__ == "__main__":
    main()
