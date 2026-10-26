#!/usr/bin/env python3
"""Génère un courrier Word au format du Dépt Contrats et Conseils.

Usage :
    python3 courrier_docx.py lettre.json -o courriers/lettre.docx [--modele entete.docx]

`lettre.json` contient un courrier (objet JSON) ou plusieurs (liste) ; chaque
courrier commence sur une nouvelle page. Avec --modele, l'en-tête (logo) et le
pied de page du document Word fourni sont conservés et son contenu est remplacé.

Mise en forme dans les textes : **gras**, __souligné__ (combinables).
Lignes du corps : "" = ligne vide, "- texte" = tiret, "1. texte" = point numéroté.
Voir exemple.json pour tous les champs.
"""
import argparse
import json
import re
import sys
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Arial"
SIZE = Pt(12)
INDENT_DESTINATAIRE = Cm(8.5)
INDENT_SIGNATURE = Cm(9)
INDENT_LISTE = Cm(1)
INDENT_SUITE_LIGNE = Cm(1.6)
MARKUP = re.compile(r"(\*\*.+?\*\*|__.+?__)")
LISTE = re.compile(r"^(-|\d+[.)])\s+(.*)$")

EXPEDITEUR = ["BRANCHE CARBURANTS", "DIRECTION JURIDIQUE", "Dépt Contrats et Conseils"]
LIEU = "Dar El Beida, le"
SIGNATURE = ["LA DIRECTRICE JURIDIQUE", "[SIGNATAIRE]"]


def segments(text, bold=False, underline=False):
    """Découpe un texte balisé en (morceau, gras, souligné)."""
    pos = 0
    for m in MARKUP.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], bold, underline
        token = m.group(0)
        if token.startswith("**"):
            yield from segments(token[2:-2], True, underline)
        else:
            yield from segments(token[2:-2], bold, True)
        pos = m.end()
    if pos < len(text):
        yield text[pos:], bold, underline


def add_runs(p, text, bold=False):
    for seg, b, u in segments(text):
        run = p.add_run(seg)
        run.font.name = FONT
        run.font.size = SIZE
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)
        run.bold = bold or b
        run.underline = u


def para(doc, text="", bold=False, left=None, first=None, after=0, line=1.0,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    p.alignment = align
    if left is not None:
        pf.left_indent = left
    if first is not None:
        pf.first_line_indent = first
    add_runs(p, text, bold)
    return p


def preparer_document(modele):
    if modele:
        doc = Document(modele)
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    else:
        doc = Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin, s.right_margin = Cm(2.5), Cm(1.5)
        s.bottom_margin, s.left_margin = Cm(1.25), Cm(2.25)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = SIZE
    return doc


def ligne_etiquette(doc, label, valeur):
    """'Label : valeur' ; une liste de valeurs s'aligne sous la première."""
    valeurs = valeur if isinstance(valeur, list) else [valeur]
    p = para(doc)
    add_runs(p, f"{label} :", bold=True)
    add_runs(p, f" {valeurs[0]}")
    for v in valeurs[1:]:
        para(doc, v, left=INDENT_SUITE_LIGNE)


def ecrire_courrier(doc, c, premiere_page):
    s = doc.sections[-1]
    largeur = s.page_width - s.left_margin - s.right_margin

    expediteur = c.get("expediteur", EXPEDITEUR)
    p = para(doc)
    p.paragraph_format.tab_stops.add_tab_stop(largeur, WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.page_break_before = not premiere_page
    add_runs(p, expediteur[0], bold=True)
    add_runs(p, "\t" + c.get("lieu", LIEU) + " ")
    for ligne in expediteur[1:]:
        para(doc, ligne, bold=True)
    para(doc, c.get("ref", f"Réf : N°                    /{c.get('annee', date.today().year)}"),
         bold=True)
    para(doc)

    for ligne in c["destinataire"]:
        para(doc, ligne, bold=True, left=INDENT_DESTINATAIRE)
    para(doc)

    ligne_etiquette(doc, "Objet", c["objet"])
    for label, valeur in c.get("lignes", []):
        ligne_etiquette(doc, label, valeur)
    para(doc)

    if c.get("appel"):
        para(doc, c["appel"], after=10)
    corps = c.get("corps", [])
    for i, ligne in enumerate(corps):
        m = LISTE.match(ligne)
        if not ligne.strip():
            para(doc)
        elif m:
            fin_liste = i + 1 == len(corps) or not LISTE.match(corps[i + 1])
            p = para(doc, after=10 if fin_liste else 4, line=1.15,
                     left=INDENT_LISTE, first=-Cm(0.6))
            p.paragraph_format.tab_stops.add_tab_stop(INDENT_LISTE)
            add_runs(p, m.group(1) + "\t" + m.group(2))
        else:
            para(doc, ligne, after=10, line=1.15)

    para(doc, c.get("cloture", "Salutations"), after=10)
    para(doc)
    for ligne in c.get("signature", SIGNATURE):
        para(doc, ligne, bold=True, left=INDENT_SIGNATURE)

    copies = c.get("copie", [])
    if copies:
        para(doc)
        ligne_etiquette(doc, "Copie", copies)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("entree", help="fichier JSON (ou - pour l'entrée standard)")
    ap.add_argument("-o", "--sortie", required=True, help="fichier .docx à créer")
    ap.add_argument("--modele", help=".docx dont l'en-tête et le pied de page sont repris")
    args = ap.parse_args()

    source = sys.stdin if args.entree == "-" else open(args.entree, encoding="utf-8")
    with source:
        data = json.load(source)
    courriers = data if isinstance(data, list) else [data]

    doc = preparer_document(args.modele)
    for i, c in enumerate(courriers):
        ecrire_courrier(doc, c, premiere_page=(i == 0))
    doc.save(args.sortie)
    print(f"{len(courriers)} courrier(s) -> {args.sortie}")


if __name__ == "__main__":
    main()
